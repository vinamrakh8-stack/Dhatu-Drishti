"""Downloadable DOCX/PDF rendering of a generated analysis."""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

UNAVAILABLE = "Data not available"


def _as_list(value) -> list[str]:
    if isinstance(value, list):
        return [str(v) for v in value if str(v).strip()]
    if isinstance(value, str) and value.strip():
        return [value]
    return []


def _sections(report: dict) -> list[tuple[str, Any]]:
    """Ordered (heading, body) pairs; body is a str or list of str."""
    prediction = report.get("short_term_prediction") or {}
    trend = report.get("production_trend") or {}
    return [
        ("Mine", report.get("mine") or UNAVAILABLE),
        ("Overview", report.get("overview") or UNAVAILABLE),
        ("Current Operational Condition", report.get("operational_condition") or UNAVAILABLE),
        ("Historical Production Trend",
         f"{trend.get('direction', UNAVAILABLE)} — {trend.get('evidence', UNAVAILABLE)}"),
        ("Reserve Status", report.get("reserve_status") or UNAVAILABLE),
        ("Equipment / Downtime Risk", report.get("equipment_risk") or UNAVAILABLE),
        ("Weather Impact", report.get("weather_impact") or UNAVAILABLE),
        ("Geographic / Environmental Risk", report.get("geographic_risk") or UNAVAILABLE),
        ("Current Risk Level", report.get("risk_level") or UNAVAILABLE),
        ("Major Risk Factors", _as_list(report.get("risk_factors")) or [UNAVAILABLE]),
        ("Recommended Operational Actions",
         _as_list(report.get("recommended_actions")) or [UNAVAILABLE]),
        ("Short-Term Production Outlook",
         f"Direction: {prediction.get('direction', UNAVAILABLE)} | "
         f"Risk: {prediction.get('risk', UNAVAILABLE)} | "
         f"Confidence: {prediction.get('confidence', UNAVAILABLE)}\n"
         f"{prediction.get('forecast', UNAVAILABLE)}"),
        ("Data Limitations", _as_list(report.get("data_limitations")) or [UNAVAILABLE]),
    ]


def _caveat() -> str:
    return ("Predictions are model-generated estimates derived from the supplied datasets. "
            "They are not guaranteed future observations and must be validated by on-site "
            "assessment before operational decisions are taken.")


def build_text(report: dict, mine_name: str, generated_at: str) -> str:
    """Plain-text rendering of the same sections, stored alongside the JSON.

    Used for the admin report view and as a fallback when ``report_data`` is
    missing.  Reuses :func:`_sections` so the three renderers never drift.
    """
    lines = [
        f"Mining Intelligence Report — {mine_name}",
        f"Generated: {generated_at}",
        f"Model: {report.get('model', UNAVAILABLE)}",
        _caveat(),
        "",
    ]
    for heading, body in _sections(report):
        lines.append(heading)
        if isinstance(body, list):
            lines.extend(f"  - {item}" for item in body)
        else:
            lines.extend(f"  {part}" for part in str(body).splitlines())
        lines.append("")
    return "\n".join(lines).strip()


def build_docx(report: dict, mine_name: str, generated_at: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(f"Mining Intelligence Report — {mine_name}", level=0)
    document.add_paragraph(f"Generated: {generated_at}")
    document.add_paragraph(f"Model: {report.get('model', UNAVAILABLE)}")
    document.add_paragraph(_caveat())

    for heading, body in _sections(report):
        document.add_heading(heading, level=1)
        if isinstance(body, list):
            for item in body:
                document.add_paragraph(str(item), style="List Bullet")
        else:
            document.add_paragraph(str(body))

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf(report: dict, mine_name: str, generated_at: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=f"Mining Intelligence Report - {mine_name}",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title2", parent=styles["Title"], fontSize=17, leading=21)
    heading_style = ParagraphStyle(
        "Heading2Custom", parent=styles["Heading2"], fontSize=12.5, leading=15, spaceBefore=10
    )
    body_style = ParagraphStyle("Body2", parent=styles["BodyText"], fontSize=9.5, leading=13.5)
    small_style = ParagraphStyle("Small", parent=body_style, fontSize=8.5, textColor="#444444")

    def esc(text: str) -> str:
        return (str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                .replace("\n", "<br/>"))

    story = [
        Paragraph(f"Mining Intelligence Report — {esc(mine_name)}", title_style),
        Paragraph(f"Generated: {esc(generated_at)}", small_style),
        Paragraph(f"Model: {esc(report.get('model', UNAVAILABLE))}", small_style),
        Paragraph(esc(_caveat()), small_style),
        Spacer(1, 6),
    ]

    for heading, body in _sections(report):
        story.append(Paragraph(esc(heading), heading_style))
        if isinstance(body, list):
            for item in body:
                story.append(Paragraph(f"• {esc(item)}", body_style))
        else:
            story.append(Paragraph(esc(body), body_style))

    doc.build(story)
    return buffer.getvalue()


def safe_filename(mine_name: str, extension: str) -> str:
    slug = "".join(ch if ch.isalnum() else "-" for ch in mine_name.lower())
    slug = "-".join(part for part in slug.split("-") if part)[:60] or "mine"
    stamp = datetime.now().strftime("%Y%m%d")
    return f"dhatu-drishti-{slug}-{stamp}.{extension}"
